from pathlib import Path

from openpyxl import load_workbook
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
EXPECTED = {
    "concept-context-grid.xlsx": ["Concept Grid", "Instructions"],
    "source-family-inventory.xlsx": ["Source Families", "Instructions"],
    "search-ai-audit-log.xlsx": ["Search Log", "AI Audit", "Instructions"],
    "evidence-map-stopping-rule.xlsx": [
        "Evidence Map",
        "Gap Analysis",
        "Collection Plan",
        "Stopping Rule",
        "Instructions",
    ],
    "literature-concept-vocabulary-map.xlsx": ["Concepts", "Vocabulary", "Instructions"],
    "literature-search-ai-log.xlsx": [
        "Search Routes",
        "Citation Chaining",
        "AI Audit",
        "Instructions",
    ],
    "literature-source-evaluation.xlsx": [
        "Source Register",
        "Claim Register",
        "Instructions",
    ],
    "literature-synthesis-stopping-rule.xlsx": [
        "Synthesis Matrix",
        "Coverage Audit",
        "Memo Plan",
        "Stopping Rule",
        "Instructions",
    ],
}


def test_research_workbooks_have_required_sheets() -> None:
    for filename, sheet_names in EXPECTED.items():
        workbook = load_workbook(ROOT / "templates" / filename, read_only=False)
        assert workbook.sheetnames == sheet_names
        for worksheet in workbook.worksheets:
            assert worksheet.max_row >= 2
            assert worksheet.max_column >= 2
            assert worksheet["A1"].value


def test_workbooks_include_license_and_empty_reader_row() -> None:
    for filename in EXPECTED:
        workbook = load_workbook(ROOT / "templates" / filename, data_only=False)
        values = [cell.value for row in workbook["Instructions"].iter_rows() for cell in row]
        assert any("CC BY 4.0" in str(value) for value in values)
        first = workbook[workbook.sheetnames[0]]
        assert all(first.cell(row=3, column=column).value is None for column in range(1, first.max_column + 1))


def test_question_framing_word_template() -> None:
    document = Document(ROOT / "templates/question-framing.docx")
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    text = "\n".join(paragraph_text + table_text)
    assert "Question-Framing Worksheet" in text
    assert "This project does not attempt to" in text
    assert "CC BY 4.0" in text


def test_literature_review_protocol_word_template() -> None:
    document = Document(ROOT / "templates/literature-review-protocol.docx")
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    text = "\n".join(paragraph_text + table_text)
    for expected in (
        "Literature Review Protocol",
        "Review purpose",
        "Inclusion logic",
        "Exclusion logic",
        "Reopening trigger",
        "CC BY 4.0",
    ):
        assert expected in text
