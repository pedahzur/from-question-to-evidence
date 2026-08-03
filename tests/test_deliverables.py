from pathlib import Path
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "build-manuscript/From-Question-to-Evidence.docx"
MARKDOWN = ROOT / "build-manuscript/From-Question-to-Evidence.md"


def test_discussion_draft_has_expected_structure() -> None:
    document = Document(DOCX)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading 1")
    ]
    assert "Ami Pedahzur" in text
    assert "Jonathan Grossman" in text
    assert len(headings) >= 19
    assert len(text.split()) >= 15_000
    assert "Literature as Evidence" in text
    assert "Synthesis Matrix" in text
    assert "Stopping Rule" in text


def test_discussion_draft_contains_live_notes() -> None:
    with ZipFile(DOCX) as archive:
        names = set(archive.namelist())
        assert "word/footnotes.xml" in names
        footnotes = archive.read("word/footnotes.xml").decode("utf-8")
    assert "NIST" in footnotes
    assert "Global Terrorism Database" in footnotes


def test_complete_markdown_export_contains_the_full_current_book() -> None:
    text = MARKDOWN.read_text(encoding="utf-8")

    assert "Building Event Databases with AI" in text
    assert "Bayesian evidence table" in text
    assert "Historical Sources as Evidence" in text
    assert "Historical Evidence Chain Register" in text
    assert "AI and Research Integrity" in text
    assert "References" in text
    assert len(text.split()) >= 23_000
